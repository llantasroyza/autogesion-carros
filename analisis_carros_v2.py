import streamlit as st
import pandas as pd
import gspread
from google.oauth2.service_account import Credentials
import plotly.express as px
from datetime import date, datetime, timedelta
import base64, io, urllib.parse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Gallery Motors by UR", page_icon="🏎️", layout="wide", initial_sidebar_state="collapsed")

SCOPES = ["https://www.googleapis.com/auth/spreadsheets","https://www.googleapis.com/auth/drive"]
MONEDAS = ["CRC","USD"]
TIPOS_GASTO_VEHICULO = ["Compra","Flete","Seguro","Impuestos Importación","Bodegaje","Dekra","Mejoras Mecánicas","Mejoras Pintura","Repuestos","Marchamo","Traspaso","Gasolina y Líquidos","Comisión Vendedor","Otros"]
TIPOS_GASTO_OPERATIVO = ["Planilla","Alquiler","Agua","Luz","Teléfono","Internet","Publicidad","Útiles de Oficina","Gasolina General","Limpieza","Otros Operativos"]
ESTADOS = ["Disponible","Reservado","Vendido","Financiado","En Consignación"]
TIPOS_PROPIEDAD = ["Propio","Con Socio","Consignación"]

# ── LOGO ──
with open('/mnt/user-data/uploads/WhatsApp_Image_2026-04-15_at_8_56_21_PM.jpeg','rb') as f:
    LOGO_B64 = base64.b64encode(f.read()).decode()

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Bebas+Neue&family=DM+Sans:wght@300;400;500;600&display=swap');
:root{--negro:#0a0a0a;--blanco:#f5f0e8;--rojo:#e63946;--gris:#1a1a1a;--gris2:#2a2a2a;--oro:#c9a84c;--verde:#4caf50;--morado:#9b59b6;--azul:#3498db;}
html,body,[class*="css"]{font-family:'DM Sans',sans-serif;background-color:var(--negro);color:var(--blanco);}
.stApp{background-color:var(--negro);}
.hero{background:linear-gradient(135deg,#0a0a0a 0%,#1a1a1a 50%,#0a0a0a 100%);border-bottom:2px solid var(--rojo);padding:1.2rem 1.5rem;margin:-1rem -1rem 2rem -1rem;display:flex;align-items:center;justify-content:space-between;flex-wrap:wrap;gap:0.5rem;}
.hero-title{font-family:'Bebas Neue',cursive;font-size:clamp(1.8rem,5vw,3.5rem);letter-spacing:4px;color:var(--blanco);line-height:1;margin:0;}
.hero-title span{color:var(--rojo);}
.hero-sub{font-size:clamp(0.6rem,2vw,0.85rem);color:#888;letter-spacing:3px;text-transform:uppercase;margin-top:0.3rem;}
.metric-card{background:var(--gris);border:1px solid var(--gris2);border-left:3px solid var(--rojo);border-radius:4px;padding:1rem 1.2rem;margin-bottom:0.8rem;}
.metric-label{font-size:0.65rem;letter-spacing:2px;text-transform:uppercase;color:#888;margin-bottom:0.2rem;}
.metric-value{font-family:'Bebas Neue',cursive;font-size:1.6rem;color:var(--blanco);}
.metric-value.verde{color:var(--verde);}.metric-value.rojo{color:var(--rojo);}.metric-value.oro{color:var(--oro);}
.carro-card{background:var(--gris);border:1px solid var(--gris2);border-radius:8px;padding:1.2rem;margin-bottom:1rem;position:relative;overflow:hidden;}
.carro-card:hover{border-color:var(--rojo);}
.carro-card::before{content:'';position:absolute;top:0;left:0;width:3px;height:100%;background:var(--rojo);}
.badge{display:inline-block;padding:0.2rem 0.7rem;border-radius:20px;font-size:0.65rem;font-weight:600;letter-spacing:1px;text-transform:uppercase;margin-right:4px;}
.badge-disponible{background:#1a3a1a;color:var(--verde);}.badge-vendido{background:#3a1a1a;color:var(--rojo);}.badge-reservado{background:#3a3a1a;color:var(--oro);}.badge-financiado{background:#2a1a3a;color:var(--morado);}
.carro-nombre{font-family:'Bebas Neue',cursive;font-size:1.4rem;letter-spacing:2px;margin:0.4rem 0 0.2rem;}
.carro-precio{font-family:'Bebas Neue',cursive;font-size:1.6rem;color:var(--oro);}
.stTabs [data-baseweb="tab-list"]{background:var(--gris);border-radius:4px;padding:4px;gap:2px;flex-wrap:wrap;}
.stTabs [data-baseweb="tab"]{background:transparent;color:#888;border-radius:2px;font-size:0.68rem;letter-spacing:0.5px;text-transform:uppercase;font-weight:600;}
.stTabs [aria-selected="true"]{background:var(--rojo) !important;color:white !important;}
.stButton button{background:var(--rojo) !important;color:white !important;border:none !important;font-weight:600 !important;letter-spacing:1px !important;text-transform:uppercase !important;font-size:0.8rem !important;padding:0.5rem 1.2rem !important;border-radius:2px !important;}
.btn-sec button{background:var(--gris2) !important;border:1px solid #444 !important;}
.btn-verde button{background:#27ae60 !important;}
.btn-rojo button{background:#c0392b !important;}
.seccion-titulo{font-family:'Bebas Neue',cursive;font-size:1.6rem;letter-spacing:3px;color:var(--blanco);border-bottom:1px solid var(--gris2);padding-bottom:0.5rem;margin-bottom:1.2rem;}
.seccion-titulo span{color:var(--rojo);}
.alerta-box{background:#3a2a0a;border:1px solid #ff9800;border-radius:6px;padding:0.8rem 1rem;margin-bottom:0.5rem;font-size:0.85rem;}
hr{border-color:var(--gris2) !important;}
</style>""", unsafe_allow_html=True)

# ── CONEXIÓN ──
@st.cache_resource
def conectar():
    creds = Credentials.from_service_account_info(dict(st.secrets["gcp_service_account"]), scopes=SCOPES)
    return gspread.authorize(creds)

HOJAS_HEADERS = {
    "INVENTARIO": ["ID","Vehículo","Año","Color","Placa","VIN","Socios","Tipo Propiedad","% Grupo","% Socio","Estado","Precio Objetivo (CRC)","Precio Mínimo (CRC)","Fecha Entrada","Notas","Origen","URL Foto"],
    "GASTOS_VEHICULO": ["ID Vehículo","Fecha","Tipo Gasto","Detalle","Moneda","Monto Original","Tipo Cambio","Monto CRC","Pagado Por","Vendedor"],
    "VENTAS": ["ID Vehículo","Fecha Venta","ID Cliente","Moneda","Precio Original","Tipo Cambio","Precio CRC","Vendedor","Comisión CRC","Forma de Pago","Lleva Factura","Num Transferencia","Observaciones"],
    "CANJES": ["ID","ID Venta","Vehículo Canje","Año","Color","Placa","Valor Acordado CRC","Notas","Fecha"],
    "FINANCIAMIENTOS": ["ID","ID Vehículo","ID Cliente","Fecha Inicio","Monto Financiado","Moneda","Tasa Interés %","Plazo Meses","Cuota Mensual","Saldo Pendiente","Estado Fin","Notas"],
    "PAGOS_FINANCIAMIENTO": ["ID Financiamiento","Fecha Pago","Monto Pagado","Moneda","Capital","Interés","Saldo Restante","Método Pago","Observaciones","Recibo #"],
    "GASTOS_OPERATIVOS": ["Fecha","Tipo","Detalle","Moneda","Monto Original","Tipo Cambio","Monto CRC","Mes"],
    "CLIENTES": ["ID Cliente","Cédula","Nombre","Teléfono","Correo","Dirección","Ocupación","Empleo","Ingresos Mensuales","Referencia 1","Tel Ref 1","Referencia 2","Tel Ref 2","Notas","Fecha Registro"],
    "CONSIGNACIONES": ["ID","ID Vehículo","Dueño","Teléfono Dueño","Precio Mínimo Dueño","Comisión Gallery %","Estado","Fecha Inicio","Notas"],
    "HISTORIAL_PRECIOS": ["ID Vehículo","Fecha Cambio","Precio Anterior","Precio Nuevo","Razón"],
    "VENDEDORES": ["ID","Nombre","Teléfono","Activo"],
    "CITAS": ["ID","ID Vehículo","ID Cliente","Fecha","Hora","Vendedor Asignado","Estado","Notas","Seguimiento"],
    "LOG_ACTIVIDAD": ["ID Vehículo","Fecha","Tipo","Descripción","Usuario"],
}

@st.cache_data(ttl=120)
def cargar_todas():
    resultado = {}
    try:
        cliente = conectar()
        sheet = cliente.open("Sistema Carros")
        existentes = [ws.title for ws in sheet.worksheets()]
        for nombre, headers in HOJAS_HEADERS.items():
            try:
                if nombre not in existentes:
                    ws = sheet.add_worksheet(title=nombre, rows=1000, cols=30)
                    ws.append_row(headers)
                    resultado[nombre] = pd.DataFrame(columns=headers)
                else:
                    ws = sheet.worksheet(nombre)
                    data = ws.get_all_records()
                    resultado[nombre] = pd.DataFrame(data) if data else pd.DataFrame(columns=headers)
            except Exception as e:
                resultado[nombre] = pd.DataFrame(columns=HOJAS_HEADERS.get(nombre,[]))
    except Exception as e:
        st.error(f"Error conectando: {e}")
        for n,h in HOJAS_HEADERS.items():
            resultado[n] = pd.DataFrame(columns=h)
    return resultado

def cargar(nombre):
    return cargar_todas().get(nombre, pd.DataFrame(columns=HOJAS_HEADERS.get(nombre,[])))

def get_ws(nombre):
    return conectar().open("Sistema Carros").worksheet(nombre)

def invalidar():
    cargar_todas.clear()

def fmt(v):
    try: return f"₡{float(v):,.0f}"
    except: return "₡0"

def to_crc(monto, moneda, tc):
    try:
        m,t = float(monto), float(tc) if tc else 0
        return m*t if moneda=="USD" and t>0 else m
    except: return 0

def calcular_cuota(monto, tasa, plazo):
    try:
        m,t,p = float(monto), float(tasa)/100/12, int(plazo)
        if t==0: return m/p
        return m*(t*(1+t)**p)/((1+t)**p-1)
    except: return 0

def dias_inv(fecha):
    try: return (datetime.now()-pd.to_datetime(fecha)).days
    except: return 0

def wa_link(row):
    n=str(row.get("Vehículo","")); a=str(row.get("Año","")); p=fmt(row.get("Precio Objetivo (CRC)",0))
    msg=f"🚗 *{n} {a}*\n💰 Precio: {p}\n\nGallery Motors by UR"
    return f"https://wa.me/?text={urllib.parse.quote(msg)}"

def calcular_inv(df_inv, df_gv, df_ven):
    if df_inv.empty: return df_inv
    df = df_inv.copy()
    df = df[df["ID"].astype(str).str.strip().str.len()>0]
    df = df[df["ID"].astype(str).str.lower()!="none"]
    if "Vehículo" in df.columns:
        df = df[df["Vehículo"].astype(str).str.strip().str.len()>0]
        df = df[df["Vehículo"].astype(str).str.lower()!="none"]
    df["ID"] = df["ID"].astype(str)
    if not df_gv.empty and "ID Vehículo" in df_gv.columns:
        g = df_gv.copy(); g["ID Vehículo"]=g["ID Vehículo"].astype(str)
        g["Monto CRC"]=pd.to_numeric(g["Monto CRC"],errors="coerce").fillna(0)
        ct=g.groupby("ID Vehículo")["Monto CRC"].sum().reset_index(); ct.columns=["ID","Costo Total"]
        df=df.merge(ct,on="ID",how="left")
        ge=g[g["Pagado Por"].astype(str).str.lower().str.contains("edgardo",na=False)]
        ce=ge.groupby("ID Vehículo")["Monto CRC"].sum().reset_index(); ce.columns=["ID","Pagado Edgardo"]
        df=df.merge(ce,on="ID",how="left")
        gs=g[g["Pagado Por"].astype(str).str.lower().str.contains("socio|luis",na=False)]
        cs=gs.groupby("ID Vehículo")["Monto CRC"].sum().reset_index(); cs.columns=["ID","Pagado Socio"]
        df=df.merge(cs,on="ID",how="left")
    else:
        df["Costo Total"]=0; df["Pagado Edgardo"]=0; df["Pagado Socio"]=0
    for c in ["Costo Total","Pagado Edgardo","Pagado Socio"]:
        df[c]=pd.to_numeric(df.get(c,0),errors="coerce").fillna(0)
    if not df_ven.empty and "ID Vehículo" in df_ven.columns:
        v=df_ven.copy(); v["ID Vehículo"]=v["ID Vehículo"].astype(str)
        v["Precio CRC"]=pd.to_numeric(v["Precio CRC"],errors="coerce").fillna(0)
        vs=v.groupby("ID Vehículo")["Precio CRC"].sum().reset_index(); vs.columns=["ID","Precio Venta"]
        df=df.merge(vs,on="ID",how="left")
    else: df["Precio Venta"]=0
    df["Precio Venta"]=pd.to_numeric(df.get("Precio Venta",0),errors="coerce").fillna(0)
    for c in ["Precio Objetivo (CRC)","Precio Mínimo (CRC)","% Grupo","% Socio"]:
        if c in df.columns: df[c]=pd.to_numeric(df[c],errors="coerce").fillna(0)
    df["Ganancia Est."]=df["Precio Objetivo (CRC)"]-df["Costo Total"]
    mask=df["Estado"].astype(str).str.lower().isin(["vendido","financiado"])
    df["Ganancia Real"]=0.0
    df.loc[mask,"Ganancia Real"]=df.loc[mask,"Precio Venta"]-df.loc[mask,"Costo Total"]
    es_socio=df["Tipo Propiedad"].astype(str).str.lower()=="con socio"
    df["Ganancia Edgardo"]=0.0; df["Ganancia Socio Real"]=0.0
    df.loc[es_socio&mask,"Ganancia Edgardo"]=df.loc[es_socio&mask,"Ganancia Real"]*0.5
    df.loc[es_socio&mask,"Ganancia Socio Real"]=df.loc[es_socio&mask,"Ganancia Real"]*0.5
    df.loc[~es_socio&mask,"Ganancia Edgardo"]=df.loc[~es_socio&mask,"Ganancia Real"]*df.loc[~es_socio&mask,"% Grupo"]
    df["Deuda Entre Socios"]=0.0
    df.loc[es_socio,"Deuda Entre Socios"]=df.loc[es_socio,"Pagado Edgardo"]-df.loc[es_socio,"Pagado Socio"]
    if "Fecha Entrada" in df.columns:
        df["Días Inventario"]=df["Fecha Entrada"].apply(dias_inv)
    else: df["Días Inventario"]=0
    return df

# ── ESTADO ──
for k,v in [("panel",False),("auth",False),("canje_pendiente",False),("ir_a_canje",False),("venta_id",""),("modo_venta",False)]:
    if k not in st.session_state: st.session_state[k]=v

# ── HEADER ──
st.markdown(f"""<div class="hero">
  <div style="display:flex;align-items:center;gap:1rem;flex-wrap:wrap;">
    <img src="data:image/jpeg;base64,{LOGO_B64}" style="height:60px;width:auto;border-radius:4px;background:white;padding:4px;">
    <div><div class="hero-title">Gallery <span>Motors</span></div><div class="hero-sub">by UR · Sistema Integral de Gestión</div></div>
  </div>
  <div style="color:#e63946;font-size:0.75rem;font-weight:600;letter-spacing:1px;">● EN LÍNEA</div>
</div>""", unsafe_allow_html=True)

# ── NAV ──
c1,c2,c3 = st.columns([5,2,2])
with c2:
    if not st.session_state["panel"]:
        if st.button("🔐 Panel Privado"): st.session_state["panel"]=True; st.rerun()
    else:
        st.markdown('<div class="btn-sec">',unsafe_allow_html=True)
        if st.button("🏠 Inicio"): st.session_state["panel"]=False; st.session_state["auth"]=False; st.rerun()
        st.markdown('</div>',unsafe_allow_html=True)
with c3:
    if st.session_state.get("auth") and st.session_state.get("panel"):
        st.markdown('<div class="btn-sec">',unsafe_allow_html=True)
        if st.button("🚪 Cerrar sesión"):
            for k in ["auth","panel","canje_pendiente","ir_a_canje","modo_venta"]:
                st.session_state[k]=False
            st.rerun()
        st.markdown('</div>',unsafe_allow_html=True)

# ── CARGAR DATOS ──
datos = cargar_todas()
df_inv=datos["INVENTARIO"]; df_gv=datos["GASTOS_VEHICULO"]; df_ven=datos["VENTAS"]
df_canjes=datos["CANJES"]; df_fin=datos["FINANCIAMIENTOS"]; df_pagos=datos["PAGOS_FINANCIAMIENTO"]
df_op=datos["GASTOS_OPERATIVOS"]; df_cli=datos["CLIENTES"]; df_cons=datos["CONSIGNACIONES"]
df_hist=datos["HISTORIAL_PRECIOS"]; df_vend=datos["VENDEDORES"]; df_citas=datos["CITAS"]; df_log=datos["LOG_ACTIVIDAD"]
df_calc=calcular_inv(df_inv, df_gv, df_ven)

def vlabel(x):
    r=df_inv[df_inv["ID"].astype(str)==str(x)]
    return f"ID {x} - {r['Vehículo'].values[0]}" if not r.empty and "Vehículo" in r.columns else f"ID {x}"

def clabel(x):
    if not df_cli.empty and "ID Cliente" in df_cli.columns:
        r=df_cli[df_cli["ID Cliente"].astype(str)==str(x)]
        if not r.empty: return f"{r['Nombre'].values[0]}"
    return f"ID {x}"

ids_v = df_inv["ID"].astype(str).tolist() if not df_inv.empty and "ID" in df_inv.columns else []
ids_cli = df_cli["ID Cliente"].astype(str).tolist() if not df_cli.empty and "ID Cliente" in df_cli.columns else []
vend_lista = df_vend["Nombre"].tolist() if not df_vend.empty and "Nombre" in df_vend.columns else []

# ════════════════════════════════
# CATÁLOGO PÚBLICO
# ════════════════════════════════
if not st.session_state["panel"]:
    disp=df_calc[df_calc["Estado"].astype(str).str.lower()=="disponible"] if not df_calc.empty else pd.DataFrame()
    res=df_calc[df_calc["Estado"].astype(str).str.lower()=="reservado"] if not df_calc.empty else pd.DataFrame()
    cons_c=df_calc[df_calc["Estado"].astype(str).str.lower()=="en consignación"] if not df_calc.empty else pd.DataFrame()
    c1,c2,c3,c4=st.columns(4)
    inv_act=disp["Costo Total"].sum() if not disp.empty and "Costo Total" in disp.columns else 0
    for col,lbl,val,cls in [(c1,"Disponibles",len(disp),"verde"),(c2,"Reservados",len(res),"oro"),(c3,"Consignación",len(cons_c),"azul"),(c4,"Inversión Activa",fmt(inv_act),"oro")]:
        with col: st.markdown(f'<div class="metric-card"><div class="metric-label">{lbl}</div><div class="metric-value {cls}">{val}</div></div>',unsafe_allow_html=True)
    st.markdown('<div class="seccion-titulo">CATÁLOGO <span>ACTUAL</span></div>',unsafe_allow_html=True)
    f1,f2=st.columns([2,3])
    with f1: filtro=st.selectbox("Estado",["Todos","Disponible","Reservado","En Consignación"])
    with f2: buscar=st.text_input("🔍 Buscar")
    df_cat=df_calc[~df_calc["Estado"].astype(str).str.lower().isin(["vendido","financiado"])].copy() if not df_calc.empty else pd.DataFrame()
    if not df_cat.empty:
        if filtro!="Todos": df_cat=df_cat[df_cat["Estado"].astype(str).str.lower()==filtro.lower()]
        if buscar: df_cat=df_cat[df_cat["Vehículo"].astype(str).str.contains(buscar,case=False,na=False)]
    if df_cat.empty:
        st.markdown('<div style="text-align:center;padding:4rem;color:#444;"><div style="font-family:\'Bebas Neue\',cursive;font-size:3rem;">SIN VEHÍCULOS</div></div>',unsafe_allow_html=True)
    else:
        cols3=st.columns(3)
        for i,(_,row) in enumerate(df_cat.iterrows()):
            est=str(row.get("Estado","")).lower()
            bc=f"badge-{est.split()[0]}"
            nombre=str(row.get('Vehículo',''))
            año=str(row.get('Año',''))
            color=str(row.get('Color',''))
            placa=str(row.get('Placa',''))
            precio=fmt(row.get('Precio Objetivo (CRC)',0))
            notas=str(row.get('Notas','') or '')
            sub=' · '.join(filter(None,[año,color,f'Placa:{placa}' if placa else '']))
            wa=wa_link(row)
            origen=str(row.get('Origen',''))
            tag_c=' <span style="background:#2a1a3a;color:#9b59b6;padding:0.1rem 0.5rem;border-radius:10px;font-size:0.65rem;">CANJE</span>' if origen=="Canje" else ""
            with cols3[i%3]:
                st.markdown(f"""<div class="carro-card">
                <span class="badge {bc}">{row.get('Estado','')}</span>{tag_c}
                <div class="carro-nombre">{nombre}</div>
                <div style="color:#888;font-size:0.8rem;margin-bottom:0.6rem;">{sub}</div>
                <div class="carro-precio">{precio}</div>
                <div style="color:#555;font-size:0.75rem;">Precio objetivo</div>
                <div style="color:#888;font-size:0.8rem;margin-top:0.5rem;">{notas}</div>
                <a href="{wa}" target="_blank" style="display:inline-block;margin-top:0.8rem;background:#25D366;color:white;padding:0.3rem 0.8rem;border-radius:4px;font-size:0.75rem;text-decoration:none;font-weight:600;">📲 WhatsApp</a>
                </div>""",unsafe_allow_html=True)

# ════════════════════════════════
# PANEL PRIVADO
# ════════════════════════════════
else:
    if not st.session_state["auth"]:
        st.markdown("<div style='max-width:400px;margin:3rem auto;text-align:center;padding:3rem;background:#1a1a1a;border:1px solid #2a2a2a;border-radius:8px;'><div style='color:white;font-size:2rem;letter-spacing:3px;margin-bottom:1rem;font-weight:bold;'>🔐 PANEL PRIVADO</div><div style='color:#888;font-size:0.85rem;margin-bottom:1.5rem;'>Gallery Motors by UR</div></div>",unsafe_allow_html=True)
        _,cp,_=st.columns([1,2,1])
        with cp:
            pw=st.text_input("Contraseña",type="password",key="pw_main")
            if st.button("🔑 Entrar",key="btn_entrar"):
                if pw==st.secrets.get("panel_password","carros2024"):
                    st.session_state["auth"]=True; st.rerun()
                else: st.error("❌ Contraseña incorrecta")
    else:
        st.markdown('<div class="seccion-titulo">🔐 PANEL <span>PRIVADO</span></div>',unsafe_allow_html=True)

        # Alertas
        if not df_calc.empty and "Días Inventario" in df_calc.columns:
            alertas=df_calc[(df_calc["Días Inventario"]>60)&(df_calc["Estado"].astype(str).str.lower()=="disponible")]
            for _,a in alertas.iterrows():
                cls="🔴" if a["Días Inventario"]>90 else "🟡"
                st.markdown(f'<div class="alerta-box">{cls} <b>{a["Vehículo"]}</b> lleva <b>{int(a["Días Inventario"])} días</b> sin venderse.</div>',unsafe_allow_html=True)
        if not df_calc.empty and "Precio Mínimo (CRC)" in df_calc.columns:
            df_calc["Precio Mínimo (CRC)"]=pd.to_numeric(df_calc["Precio Mínimo (CRC)"],errors="coerce").fillna(0)
            pb=df_calc[(df_calc["Precio Mínimo (CRC)"]>0)&(df_calc["Precio Objetivo (CRC)"]<df_calc["Precio Mínimo (CRC)"])&(df_calc["Estado"].astype(str).str.lower()=="disponible")]
            for _,p in pb.iterrows():
                st.markdown(f'<div class="alerta-box">⚠️ <b>{p["Vehículo"]}</b>: Precio objetivo {fmt(p["Precio Objetivo (CRC)"])} está bajo el mínimo {fmt(p["Precio Mínimo (CRC)"])}.</div>',unsafe_allow_html=True)

        TABS=["📊 Dashboard","🚗 Vehículos","💰 Gastos Vehículo","🏆 Ventas","💳 Financiamientos","💵 Pagos/Recibos","🏢 Gastos Operativos","👥 Clientes","🤝 Socios","📦 Consignación","📈 Informes","➕ Nuevo Vehículo","⚙️ Config","📥 Importar","📤 Exportar","📧 Email","📅 Citas","📋 Log"]
        tabs=st.tabs(TABS)

        # ══ DASHBOARD ══
        with tabs[0]:
            if df_calc.empty: st.info("Sin datos.")
            else:
                gan_e=df_calc["Ganancia Edgardo"].sum() if "Ganancia Edgardo" in df_calc.columns else 0
                gan_s=df_calc["Ganancia Socio Real"].sum() if "Ganancia Socio Real" in df_calc.columns else 0
                total_inv=df_calc["Costo Total"].sum() if "Costo Total" in df_calc.columns else 0
                gastos_op=pd.to_numeric(df_op["Monto CRC"],errors="coerce").sum() if not df_op.empty and "Monto CRC" in df_op.columns else 0
                saldo_fin=pd.to_numeric(df_fin["Saldo Pendiente"],errors="coerce").sum() if not df_fin.empty and "Saldo Pendiente" in df_fin.columns else 0
                utilidad=gan_e-gastos_op
                disp_n=len(df_calc[df_calc["Estado"].astype(str).str.lower()=="disponible"])
                vend_n=len(df_calc[df_calc["Estado"].astype(str).str.lower().isin(["vendido","financiado"])])
                r1=st.columns(4)
                for col,(lbl,val) in zip(r1,[("Total Invertido",fmt(total_inv)),("Mi Ganancia",fmt(gan_e)),("Utilidad Neta",fmt(utilidad)),("Saldo Financiado",fmt(saldo_fin))]):
                    with col: st.metric(lbl,val)
                r2=st.columns(4)
                for col,(lbl,val) in zip(r2,[("Disponibles",disp_n),("Vendidos/Fin.",vend_n),("Gastos Op.",fmt(gastos_op)),("Clientes",len(df_cli))]):
                    with col: st.metric(lbl,val)
                st.divider()
                cg1,cg2=st.columns(2)
                with cg1:
                    ec=df_calc["Estado"].value_counts().reset_index(); ec.columns=["Estado","Cantidad"]
                    fig1=px.pie(ec,values="Cantidad",names="Estado",title="Carros por estado",hole=0.4,color_discrete_map={"Disponible":"#4caf50","Vendido":"#e63946","Reservado":"#c9a84c","Financiado":"#9b59b6"})
                    fig1.update_layout(paper_bgcolor="#1a1a1a",plot_bgcolor="#1a1a1a",font_color="#f5f0e8")
                    st.plotly_chart(fig1,use_container_width=True)
                with cg2:
                    vdf=df_calc[df_calc["Estado"].astype(str).str.lower().isin(["vendido","financiado"])]
                    if not vdf.empty:
                        fig2=px.bar(vdf,x="Vehículo",y="Ganancia Real",title="Ganancia por vehículo",color="Ganancia Real",color_continuous_scale=["#e63946","#c9a84c","#4caf50"])
                        fig2.update_layout(paper_bgcolor="#1a1a1a",plot_bgcolor="#1a1a1a",font_color="#f5f0e8",showlegend=False)
                        st.plotly_chart(fig2,use_container_width=True)

        # ══ VEHÍCULOS ══
        with tabs[1]:
            if df_calc.empty: st.info("Sin vehículos.")
            else:
                cols_m=[c for c in ["ID","Vehículo","Año","Color","Placa","Estado","Origen","Días Inventario","Costo Total","Precio Objetivo (CRC)","Precio Mínimo (CRC)","Precio Venta","Ganancia Real","Ganancia Edgardo","Ganancia Socio Real"] if c in df_calc.columns]
                st.dataframe(df_calc[cols_m],use_container_width=True)
                st.divider()
                st.markdown("### ✏️ Modificar vehículo")
                if ids_v:
                    id_mod=st.selectbox("Selecciona vehículo a modificar",options=ids_v,format_func=vlabel,key="mod_v")
                    row_mod=df_inv[df_inv["ID"].astype(str)==str(id_mod)]
                    if not row_mod.empty:
                        rm=row_mod.iloc[0]
                        with st.form("form_mod_v"):
                            c1,c2,c3=st.columns(3)
                            new_estado=c1.selectbox("Estado",ESTADOS,index=ESTADOS.index(str(rm.get("Estado","Disponible"))) if str(rm.get("Estado","Disponible")) in ESTADOS else 0)
                            new_precio=c2.number_input("Precio Objetivo (CRC)",value=float(rm.get("Precio Objetivo (CRC)",0) or 0),step=100000.0)
                            new_min=c3.number_input("Precio Mínimo (CRC)",value=float(rm.get("Precio Mínimo (CRC)",0) or 0),step=100000.0)
                            new_notas=st.text_area("Notas",value=str(rm.get("Notas","") or ""))
                            new_foto=st.text_input("URL Foto",value=str(rm.get("URL Foto","") or ""))
                            razon_precio=st.text_input("Razón del cambio de precio (si aplica)")
                            submit_mod=st.form_submit_button("💾 Guardar cambios")
                        if submit_mod:
                            try:
                                ws_i=get_ws("INVENTARIO"); data_i=ws_i.get_all_values()
                                headers_i=data_i[0]
                                for idx,fila in enumerate(data_i):
                                    if idx==0: continue
                                    if str(fila[0])==str(id_mod):
                                        if "Estado" in headers_i: ws_i.update_cell(idx+1,headers_i.index("Estado")+1,new_estado)
                                        if "Precio Objetivo (CRC)" in headers_i:
                                            precio_ant=rm.get("Precio Objetivo (CRC)",0)
                                            ws_i.update_cell(idx+1,headers_i.index("Precio Objetivo (CRC)")+1,new_precio)
                                            if razon_precio and float(new_precio)!=float(precio_ant or 0):
                                                get_ws("HISTORIAL_PRECIOS").append_row([str(id_mod),str(date.today()),precio_ant,new_precio,razon_precio])
                                        if "Precio Mínimo (CRC)" in headers_i: ws_i.update_cell(idx+1,headers_i.index("Precio Mínimo (CRC)")+1,new_min)
                                        if "Notas" in headers_i: ws_i.update_cell(idx+1,headers_i.index("Notas")+1,new_notas)
                                        if "URL Foto" in headers_i: ws_i.update_cell(idx+1,headers_i.index("URL Foto")+1,new_foto)
                                        break
                                invalidar(); st.success("✅ Vehículo actualizado."); st.rerun()
                            except Exception as e: st.error(f"Error: {e}")

                st.divider()
                st.markdown("### 💰 Gastos del vehículo")
                if ids_v:
                    id_sel=st.selectbox("Vehículo",options=ids_v,format_func=vlabel,key="sel_gv")
                    if not df_gv.empty and "ID Vehículo" in df_gv.columns:
                        gv_f=df_gv[df_gv["ID Vehículo"].astype(str)==str(id_sel)]
                        if not gv_f.empty:
                            st.dataframe(gv_f,use_container_width=True)
                            st.markdown(f"**Total: {fmt(pd.to_numeric(gv_f['Monto CRC'],errors='coerce').sum())}**")
                        else: st.info("Sin gastos.")

        # ══ GASTOS VEHÍCULO ══
        with tabs[2]:
            st.markdown("### 💰 Registrar gasto")
            if not ids_v:
                st.warning("Primero agrega vehículos.")
            else:
                with st.form("form_gv"):
                    c1,c2=st.columns(2)
                    id_g=c1.selectbox("Vehículo",options=ids_v,format_func=vlabel)
                    fecha_g=c2.date_input("Fecha",value=date.today())
                    c3,c4,c5=st.columns(3)
                    tipo_g=c3.selectbox("Tipo",TIPOS_GASTO_VEHICULO)
                    moneda_g=c4.selectbox("Moneda",MONEDAS)
                    pagado_por=c5.selectbox("Pagado por",["Edgardo","Socio","Ambos"])
                    detalle=st.text_input("Detalle del gasto")
                    c6,c7=st.columns(2)
                    monto_g=c6.number_input("Monto original",min_value=0.0,step=1000.0)
                    tc_g=c7.number_input("Tipo de cambio (solo si USD)",min_value=0.0,step=1.0,value=0.0)
                    if moneda_g=="USD" and tc_g>0:
                        st.info(f"💱 Monto en CRC: {fmt(monto_g*tc_g)}")
                    vendedor_g=st.text_input("Vendedor (solo si es comisión)")
                    submit_g=st.form_submit_button("💾 Guardar gasto")
                if submit_g:
                    monto_crc=to_crc(monto_g,moneda_g,tc_g)
                    try:
                        get_ws("GASTOS_VEHICULO").append_row([str(id_g),str(fecha_g),tipo_g,detalle,moneda_g,monto_g,tc_g if moneda_g=="USD" else 0,monto_crc,pagado_por,vendedor_g])
                        invalidar(); st.success("✅ Gasto guardado."); st.rerun()
                    except Exception as e: st.error(f"Error: {e}")

            if not df_gv.empty:
                st.divider()
                st.markdown("### ✏️ Eliminar/Editar gasto")
                st.dataframe(df_gv,use_container_width=True)
                st.markdown("**Para eliminar un gasto, ve directamente a Google Sheets y borra la fila.**")

        # ══ VENTAS ══
        with tabs[3]:
            # Modo cierre de venta completo
            st.markdown("### 🏆 Registrar venta")
            st.info("Complete todos los detalles de la venta antes de guardar.")
            if not ids_v:
                st.warning("Primero agrega vehículos.")
            else:
                with st.form("form_venta_completa"):
                    st.markdown("#### 🚗 Datos del vehículo")
                    c1,c2=st.columns(2)
                    id_v=c1.selectbox("Vehículo",options=ids_v,format_func=vlabel)
                    fecha_v=c2.date_input("Fecha venta",value=date.today())

                    st.markdown("#### 👤 Cliente")
                    id_cli_v=st.selectbox("Cliente",options=["Sin cliente"]+ids_cli,format_func=lambda x: clabel(x) if x!="Sin cliente" else "Sin cliente")

                    st.markdown("#### 💰 Precio de venta")
                    c3,c4,c5=st.columns(3)
                    moneda_v=c3.selectbox("Moneda",MONEDAS)
                    precio_v=c4.number_input("Precio venta",min_value=0.0,step=1000.0)
                    tc_v=c5.number_input("Tipo cambio (si USD)",min_value=0.0,step=1.0,value=0.0)
                    if moneda_v=="USD" and tc_v>0:
                        st.info(f"💱 En CRC: {fmt(precio_v*tc_v)}")

                    st.markdown("#### 💳 Forma de pago")
                    forma_pago=st.selectbox("Forma de pago principal",["Efectivo","Transferencia","Cheque","Canje","Mixto (Efectivo+Transferencia)","Mixto (Efectivo+Canje)","Mixto (Transferencia+Canje)","Financiado"])
                    num_trans=""
                    if "Transferencia" in forma_pago:
                        num_trans=st.text_input("Número de transferencia")
                    incluye_canje="Canje" in forma_pago

                    st.markdown("#### 👥 Vendedor y comisión")
                    c6,c7=st.columns(2)
                    vendedor_v=c6.selectbox("Vendedor",options=["Ninguno"]+vend_lista)
                    comision_v=c7.number_input("Comisión vendedor (CRC)",min_value=0.0,step=5000.0)

                    st.markdown("#### 🧾 Facturación")
                    c8,c9=st.columns(2)
                    lleva_factura=c8.selectbox("¿Lleva factura?",["No","Sí"])
                    rfc=""
                    if lleva_factura=="Sí": rfc=c9.text_input("Datos para factura")
                    observ=st.text_area("Observaciones")
                    submit_v=st.form_submit_button("💾 Registrar venta completa")

                if submit_v:
                    precio_crc=to_crc(precio_v,moneda_v,tc_v)
                    try:
                        get_ws("VENTAS").append_row([str(id_v),str(fecha_v),str(id_cli_v),moneda_v,precio_v,tc_v if moneda_v=="USD" else 0,precio_crc,vendedor_v,comision_v,forma_pago,lleva_factura,num_trans,observ])
                        ws_i=get_ws("INVENTARIO"); data_i=ws_i.get_all_values()
                        nuevo_estado="Financiado" if forma_pago=="Financiado" else "Vendido"
                        for idx,fila in enumerate(data_i):
                            if idx==0: continue
                            if str(fila[0])==str(id_v):
                                col_est=data_i[0].index("Estado")+1
                                ws_i.update_cell(idx+1,col_est,nuevo_estado); break
                        invalidar()
                        st.success(f"✅ Venta registrada como {nuevo_estado}.")
                        if incluye_canje:
                            st.session_state["canje_pendiente"]=True
                            st.session_state["venta_id"]=str(id_v)
                        st.rerun()
                    except Exception as e: st.error(f"Error: {e}")

                # Flujo de canje automático
                if st.session_state.get("canje_pendiente"):
                    st.warning(f"⚠️ La venta incluye CANJE. Registra el vehículo recibido:")
                    with st.form("form_canje_rapido"):
                        st.markdown("#### 🔄 Vehículo recibido como canje")
                        c1,c2,c3=st.columns(3)
                        canje_vehiculo=c1.text_input("Marca y modelo")
                        canje_año=c2.number_input("Año",min_value=1950,max_value=2030,step=1,value=2020)
                        canje_color=c3.text_input("Color")
                        c4,c5=st.columns(2)
                        canje_placa=c4.text_input("Placa")
                        canje_valor=c5.number_input("Valor acordado (CRC)",min_value=0.0,step=100000.0)
                        canje_notas=st.text_area("Notas del canje")
                        submit_canje=st.form_submit_button("✅ Registrar canje y agregar al inventario")
                    if submit_canje:
                        try:
                            id_nuevo=len(df_inv)+1 if not df_inv.empty else 1
                            get_ws("INVENTARIO").append_row([id_nuevo,canje_vehiculo,int(canje_año),canje_color,canje_placa,"","Edgardo","Propio",1,0,"Disponible",canje_valor,canje_valor,str(date.today()),canje_notas,"Canje",""])
                            n_canjes=len(df_canjes)+1 if not df_canjes.empty else 1
                            get_ws("CANJES").append_row([n_canjes,st.session_state.get("venta_id",""),canje_vehiculo,canje_año,canje_color,canje_placa,canje_valor,canje_notas,str(date.today())])
                            st.session_state["canje_pendiente"]=False
                            invalidar(); st.success(f"✅ Canje registrado. {canje_vehiculo} agregado al inventario."); st.rerun()
                        except Exception as e: st.error(f"Error: {e}")
                    ca,cb=st.columns(2)
                    with cb:
                        st.markdown('<div class="btn-sec">',unsafe_allow_html=True)
                        if st.button("❌ Registrar canje después"):
                            st.session_state["canje_pendiente"]=False; st.rerun()
                        st.markdown('</div>',unsafe_allow_html=True)

                if not df_ven.empty:
                    st.divider(); st.dataframe(df_ven,use_container_width=True)

        # ══ FINANCIAMIENTOS ══
        with tabs[4]:
            st.markdown("### 💳 Registrar financiamiento")
            if not ids_v:
                st.warning("Primero agrega vehículos.")
            else:
                with st.form("form_fin"):
                    c1,c2=st.columns(2)
                    id_fin_v=c1.selectbox("Vehículo",options=ids_v,format_func=vlabel)
                    id_fin_cli=c2.selectbox("Cliente",options=["Sin cliente"]+ids_cli,format_func=lambda x: clabel(x) if x!="Sin cliente" else "Sin cliente")
                    fecha_ini=st.date_input("Fecha inicio",value=date.today())
                    c3,c4,c5=st.columns(3)
                    monto_fin=c3.number_input("Monto financiado",min_value=0.0,step=100000.0)
                    moneda_fin=c4.selectbox("Moneda",MONEDAS)
                    tasa_fin=c5.number_input("Tasa % anual",min_value=0.0,max_value=100.0,step=0.5,value=18.0)
                    c6,c7=st.columns(2)
                    plazo_fin=c6.number_input("Plazo (meses)",min_value=1,max_value=120,step=1,value=12)
                    cuota=calcular_cuota(monto_fin,tasa_fin,plazo_fin)
                    c7.metric("Cuota mensual",fmt(cuota) if moneda_fin=="CRC" else f"${cuota:,.2f}")
                    notas_fin=st.text_area("Notas")
                    id_fin_n=(len(df_fin)+1) if not df_fin.empty else 1
                    submit_fin=st.form_submit_button("💾 Registrar")
                if submit_fin:
                    try:
                        get_ws("FINANCIAMIENTOS").append_row([id_fin_n,str(id_fin_v),str(id_fin_cli),str(fecha_ini),monto_fin,moneda_fin,tasa_fin,plazo_fin,round(cuota,2),monto_fin,"Activo",notas_fin])
                        invalidar(); st.success(f"✅ Financiamiento registrado. Cuota: {fmt(cuota)}"); st.rerun()
                    except Exception as e: st.error(f"Error: {e}")
            if not df_fin.empty:
                st.divider(); st.dataframe(df_fin,use_container_width=True)

        # ══ PAGOS/RECIBOS ══
        with tabs[5]:
            st.markdown("### 💵 Registrar pago")
            if df_fin.empty:
                st.info("No hay financiamientos.")
            else:
                ids_fin=df_fin["ID"].astype(str).tolist()
                def fin_lbl(x):
                    r=df_fin[df_fin["ID"].astype(str)==str(x)]
                    if r.empty: return f"ID {x}"
                    cli=clabel(r["ID Cliente"].values[0]) if "ID Cliente" in r.columns else ""
                    veh=vlabel(r["ID Vehículo"].values[0]) if "ID Vehículo" in r.columns else ""
                    return f"ID {x} - {cli} | {veh}"
                with st.form("form_pago"):
                    c1,c2=st.columns(2)
                    id_fin_sel=c1.selectbox("Financiamiento",options=ids_fin,format_func=fin_lbl)
                    fecha_pago=c2.date_input("Fecha",value=date.today())
                    fin_row=df_fin[df_fin["ID"].astype(str)==str(id_fin_sel)]
                    if not fin_row.empty:
                        fr=fin_row.iloc[0]
                        saldo=float(fr.get("Saldo Pendiente",0) or 0)
                        tasa_def=float(fr.get("Tasa Interés %",18) or 18)
                        st.info(f"**Saldo:** {fmt(saldo)} | **Tasa pactada:** {tasa_def}%")
                    else: saldo=0; tasa_def=18; fr={}
                    c3,c4=st.columns(2)
                    monto_pago=c3.number_input("Monto del pago",min_value=0.0,step=1000.0)
                    moneda_pago=c4.selectbox("Moneda",MONEDAS)
                    c5,c6=st.columns(2)
                    metodo=c5.selectbox("Método",["Efectivo","Transferencia","Cheque"])
                    tasa_pago=c6.number_input("Tasa % para este pago",min_value=0.0,max_value=100.0,step=0.5,value=tasa_def)
                    interes=saldo*(tasa_pago/100/12)
                    capital=max(0,monto_pago-interes)
                    nuevo_saldo=max(0,saldo-capital)
                    if monto_pago>0:
                        st.markdown(f"**Capital: {fmt(capital)} | Interés: {fmt(interes)} | Nuevo saldo: {fmt(nuevo_saldo)}**")
                    obs_pago=st.text_area("Observaciones")
                    num_recibo=f"GM-{datetime.now().strftime('%Y%m%d%H%M%S')}"
                    submit_pago=st.form_submit_button("💾 Registrar y generar recibo")
                if submit_pago and not fin_row.empty:
                    try:
                        get_ws("PAGOS_FINANCIAMIENTO").append_row([str(id_fin_sel),str(fecha_pago),monto_pago,moneda_pago,round(capital,2),round(interes,2),round(nuevo_saldo,2),metodo,obs_pago,num_recibo])
                        ws_fin=get_ws("FINANCIAMIENTOS"); df_fin2=ws_fin.get_all_values()
                        for idx,fila in enumerate(df_fin2):
                            if idx==0: continue
                            if str(fila[0])==str(id_fin_sel):
                                col_s=df_fin2[0].index("Saldo Pendiente")+1
                                ws_fin.update_cell(idx+1,col_s,round(nuevo_saldo,2))
                                if nuevo_saldo<=0:
                                    col_e=df_fin2[0].index("Estado Fin")+1
                                    ws_fin.update_cell(idx+1,col_e,"Pagado")
                                break
                        invalidar()
                        ms="₡" if moneda_pago=="CRC" else "$"
                        recibo=f"""<!DOCTYPE html><html><head><meta charset="UTF-8"><title>Recibo {num_recibo}</title>
                        <style>body{{font-family:Arial,sans-serif;max-width:700px;margin:40px auto;color:#111;}}
                        h1{{color:#C0392B;}} table{{width:100%;border-collapse:collapse;}} td{{padding:10px;border-bottom:1px solid #eee;}}
                        .total{{background:#C0392B;color:white;font-weight:bold;}}</style></head><body>
                        <h1>Gallery Motors by UR</h1><h2>RECIBO #{num_recibo}</h2>
                        <table>
                        <tr><td>Fecha:</td><td><b>{fecha_pago}</b></td></tr>
                        <tr><td>Vehículo:</td><td><b>{vlabel(fr.get('ID Vehículo',''))}</b></td></tr>
                        <tr><td>Método:</td><td><b>{metodo}</b></td></tr>
                        <tr><td>Capital:</td><td><b>{ms}{capital:,.2f}</b></td></tr>
                        <tr><td>Interés ({tasa_pago}% anual):</td><td><b>{ms}{interes:,.2f}</b></td></tr>
                        <tr class="total"><td>TOTAL PAGADO:</td><td>{ms}{monto_pago:,.2f}</td></tr>
                        <tr><td>Saldo restante:</td><td style="color:#C0392B;"><b>{ms}{nuevo_saldo:,.2f}</b></td></tr>
                        </table></body></html>"""
                        st.success(f"✅ Pago registrado. Recibo #{num_recibo}")
                        st.download_button("⬇️ Descargar recibo",data=recibo,file_name=f"Recibo_{num_recibo}.html",mime="text/html")
                        st.rerun()
                    except Exception as e: st.error(f"Error: {e}")
                if not df_pagos.empty:
                    st.divider(); st.dataframe(df_pagos,use_container_width=True)

        # ══ GASTOS OPERATIVOS ══
        with tabs[6]:
            st.markdown("### 🏢 Registrar gasto operativo")
            with st.form("form_op"):
                c1,c2=st.columns(2)
                fecha_op=c1.date_input("Fecha",value=date.today())
                tipo_op=c2.selectbox("Tipo",TIPOS_GASTO_OPERATIVO)
                detalle_op=st.text_input("Detalle")
                c3,c4,c5=st.columns(3)
                moneda_op=c3.selectbox("Moneda",MONEDAS)
                monto_op=c4.number_input("Monto",min_value=0.0,step=1000.0)
                tc_op=c5.number_input("Tipo cambio (si USD)",min_value=0.0,step=1.0,value=0.0)
                if moneda_op=="USD" and tc_op>0:
                    st.info(f"💱 En CRC: {fmt(monto_op*tc_op)}")
                submit_op=st.form_submit_button("💾 Guardar")
            if submit_op:
                try:
                    monto_crc_op=to_crc(monto_op,moneda_op,tc_op)
                    get_ws("GASTOS_OPERATIVOS").append_row([str(fecha_op),tipo_op,detalle_op,moneda_op,monto_op,tc_op if moneda_op=="USD" else 0,monto_crc_op,fecha_op.strftime("%Y-%m")])
                    invalidar(); st.success("✅ Gasto guardado."); st.rerun()
                except Exception as e: st.error(f"Error: {e}")
            if not df_op.empty:
                df_op["Monto CRC"]=pd.to_numeric(df_op["Monto CRC"],errors="coerce").fillna(0)
                st.metric("Total gastos operativos",fmt(df_op["Monto CRC"].sum()))
                st.dataframe(df_op,use_container_width=True)

        # ══ CLIENTES ══
        with tabs[7]:
            st.markdown("### 👥 Registrar cliente")
            with st.form("form_cli"):
                id_cli_n=(len(df_cli)+1) if not df_cli.empty else 1
                c1,c2,c3=st.columns(3)
                cedula=c1.text_input("Cédula")
                nombre_cli=c2.text_input("Nombre completo")
                tel_cli=c3.text_input("Teléfono")
                c4,c5=st.columns(2)
                correo_cli=c4.text_input("Correo")
                dir_cli=c5.text_input("Dirección")
                c6,c7,c8=st.columns(3)
                ocup=c6.text_input("Ocupación")
                empleo=c7.selectbox("Tipo empleo",["Asalariado","Independiente","Empresario","Pensionado","Otro"])
                ingresos=c8.number_input("Ingresos mensuales",min_value=0.0,step=50000.0)
                c9,c10=st.columns(2)
                ref1=c9.text_input("Referencia 1"); tel1=c10.text_input("Tel Ref 1")
                c11,c12=st.columns(2)
                ref2=c11.text_input("Referencia 2"); tel2=c12.text_input("Tel Ref 2")
                notas_cli=st.text_area("Notas")
                submit_cli=st.form_submit_button("💾 Registrar cliente")
            if submit_cli:
                try:
                    get_ws("CLIENTES").append_row([id_cli_n,cedula,nombre_cli,tel_cli,correo_cli,dir_cli,ocup,empleo,ingresos,ref1,tel1,ref2,tel2,notas_cli,str(date.today())])
                    invalidar(); st.success(f"✅ Cliente {nombre_cli} registrado."); st.rerun()
                except Exception as e: st.error(f"Error: {e}")
            if not df_cli.empty:
                st.divider()
                buscar_cli=st.text_input("🔍 Buscar cliente")
                df_cli_s=df_cli.copy()
                if buscar_cli: df_cli_s=df_cli_s[df_cli_s["Nombre"].astype(str).str.contains(buscar_cli,case=False,na=False)|df_cli_s["Cédula"].astype(str).str.contains(buscar_cli,na=False)]
                st.dataframe(df_cli_s,use_container_width=True)

        # ══ SOCIOS ══
        with tabs[8]:
            st.markdown('<div class="seccion-titulo">🤝 CONTROL DE <span>SOCIOS</span></div>',unsafe_allow_html=True)
            if df_calc.empty: st.info("Sin datos.")
            else:
                cs=df_calc[df_calc["Tipo Propiedad"].astype(str).str.lower()=="con socio"]
                if cs.empty: st.info("No hay carros con socio.")
                else:
                    for _,row in cs.iterrows():
                        deuda=float(row.get("Deuda Entre Socios",0))
                        pag_e=float(row.get("Pagado Edgardo",0))
                        pag_s=float(row.get("Pagado Socio",0))
                        gan_e=float(row.get("Ganancia Edgardo",0))
                        gan_s=float(row.get("Ganancia Socio Real",0))
                        estado=str(row.get("Estado",""))
                        if deuda>500: msg=f"💰 Socio le debe a Edgardo: {fmt(abs(deuda))}"; color="var(--verde)"
                        elif deuda<-500: msg=f"💰 Edgardo le debe al Socio: {fmt(abs(deuda))}"; color="var(--rojo)"
                        else: msg="✅ Gastos equilibrados"; color="var(--oro)"
                        st.markdown(f"""<div style="background:#1a2a3a;border:1px solid #3498db;border-radius:8px;padding:1rem;margin-bottom:0.8rem;">
                        <b style="font-size:1.1rem;">{row.get('Vehículo','')} {row.get('Año','')}</b> — <span style="color:#c9a84c;">{estado}</span>
                        <div style="display:grid;grid-template-columns:repeat(4,1fr);gap:1rem;margin-top:0.8rem;">
                        <div><div style="color:#888;font-size:0.7rem;">PAGÓ EDGARDO</div><div style="font-weight:bold;">{fmt(pag_e)}</div></div>
                        <div><div style="color:#888;font-size:0.7rem;">PAGÓ SOCIO</div><div style="font-weight:bold;">{fmt(pag_s)}</div></div>
                        <div><div style="color:#888;font-size:0.7rem;">GANANCIA EDGARDO</div><div style="font-weight:bold;color:var(--verde);">{fmt(gan_e)}</div></div>
                        <div><div style="color:#888;font-size:0.7rem;">GANANCIA SOCIO</div><div style="font-weight:bold;color:var(--verde);">{fmt(gan_s)}</div></div>
                        </div><div style="margin-top:0.8rem;color:{color};font-weight:bold;">{msg}</div>
                        </div>""",unsafe_allow_html=True)
                    st.divider()
                    tot_e=cs["Pagado Edgardo"].sum() if "Pagado Edgardo" in cs.columns else 0
                    tot_s=cs["Pagado Socio"].sum() if "Pagado Socio" in cs.columns else 0
                    gan_e_t=cs["Ganancia Edgardo"].sum() if "Ganancia Edgardo" in cs.columns else 0
                    gan_s_t=cs["Ganancia Socio Real"].sum() if "Ganancia Socio Real" in cs.columns else 0
                    deuda_t=tot_e-tot_s
                    c1,c2,c3,c4=st.columns(4)
                    with c1: st.metric("Pagado Edgardo",fmt(tot_e))
                    with c2: st.metric("Pagado Socio",fmt(tot_s))
                    with c3: st.metric("Ganancia Edgardo",fmt(gan_e_t))
                    with c4: st.metric("Ganancia Socio",fmt(gan_s_t))
                    if deuda_t>500: st.success(f"✅ Socio le debe a Edgardo: {fmt(abs(deuda_t))}")
                    elif deuda_t<-500: st.warning(f"⚠️ Edgardo le debe al Socio: {fmt(abs(deuda_t))}")
                    else: st.info("✅ Balance equilibrado.")

        # ══ CONSIGNACIÓN ══
        with tabs[9]:
            st.markdown("### 📦 Registrar consignación")
            if not ids_v:
                st.warning("Primero agrega vehículos.")
            else:
                with st.form("form_cons"):
                    c1,c2=st.columns(2)
                    id_cons=c1.selectbox("Vehículo",options=ids_v,format_func=vlabel)
                    fecha_cons=c2.date_input("Fecha inicio",value=date.today())
                    c3,c4=st.columns(2)
                    dueno=c3.text_input("Nombre del dueño")
                    tel_dueno=c4.text_input("Teléfono")
                    c5,c6=st.columns(2)
                    precio_min_d=c5.number_input("Precio mínimo dueño (CRC)",min_value=0.0,step=100000.0)
                    comision_g=c6.number_input("Comisión Gallery %",min_value=0.0,max_value=50.0,step=0.5,value=5.0)
                    notas_cons=st.text_area("Notas")
                    id_cons_n=(len(df_cons)+1) if not df_cons.empty else 1
                    submit_cons=st.form_submit_button("💾 Registrar")
                if submit_cons:
                    try:
                        get_ws("CONSIGNACIONES").append_row([id_cons_n,str(id_cons),dueno,tel_dueno,precio_min_d,comision_g,"Activa",str(fecha_cons),notas_cons])
                        ws_i=get_ws("INVENTARIO"); data_i=ws_i.get_all_values()
                        for idx,fila in enumerate(data_i):
                            if idx==0: continue
                            if str(fila[0])==str(id_cons):
                                ws_i.update_cell(idx+1,data_i[0].index("Estado")+1,"En Consignación")
                                ws_i.update_cell(idx+1,data_i[0].index("Tipo Propiedad")+1,"Consignación"); break
                        invalidar(); st.success("✅ Consignación registrada."); st.rerun()
                    except Exception as e: st.error(f"Error: {e}")
            if not df_cons.empty:
                st.divider(); st.dataframe(df_cons,use_container_width=True)

        # ══ INFORMES ══
        with tabs[10]:
            st.markdown('<div class="seccion-titulo">📈 INFORMES <span>DEL NEGOCIO</span></div>',unsafe_allow_html=True)
            if df_calc.empty: st.info("Sin datos.")
            else:
                gan_e=df_calc["Ganancia Edgardo"].sum() if "Ganancia Edgardo" in df_calc.columns else 0
                gan_s=df_calc["Ganancia Socio Real"].sum() if "Ganancia Socio Real" in df_calc.columns else 0
                total_inv=df_calc["Costo Total"].sum()
                gastos_op=pd.to_numeric(df_op["Monto CRC"],errors="coerce").sum() if not df_op.empty and "Monto CRC" in df_op.columns else 0
                comisiones=pd.to_numeric(df_ven["Comisión CRC"],errors="coerce").sum() if not df_ven.empty and "Comisión CRC" in df_ven.columns else 0
                saldo_fin=pd.to_numeric(df_fin["Saldo Pendiente"],errors="coerce").sum() if not df_fin.empty else 0
                utilidad=gan_e-gastos_op-comisiones

                st.markdown("### 💰 Resumen Financiero")
                r1=st.columns(3)
                with r1[0]:
                    st.markdown(f'<div class="metric-card"><div class="metric-label">Total Invertido</div><div class="metric-value oro">{fmt(total_inv)}</div></div>',unsafe_allow_html=True)
                    st.markdown(f'<div class="metric-card"><div class="metric-label">Ganancia Bruta</div><div class="metric-value verde">{fmt(df_calc["Ganancia Real"].sum())}</div></div>',unsafe_allow_html=True)
                with r1[1]:
                    st.markdown(f'<div class="metric-card"><div class="metric-label">Mi Ganancia (Edgardo)</div><div class="metric-value verde">{fmt(gan_e)}</div></div>',unsafe_allow_html=True)
                    st.markdown(f'<div class="metric-card"><div class="metric-label">Ganancia Socio</div><div class="metric-value oro">{fmt(gan_s)}</div></div>',unsafe_allow_html=True)
                with r1[2]:
                    st.markdown(f'<div class="metric-card"><div class="metric-label">Gastos Operativos</div><div class="metric-value rojo">{fmt(gastos_op)}</div></div>',unsafe_allow_html=True)
                    st.markdown(f'<div class="metric-card"><div class="metric-label">Utilidad Neta</div><div class="metric-value {"verde" if utilidad>=0 else "rojo"}">{fmt(utilidad)}</div></div>',unsafe_allow_html=True)

                st.divider()
                if not df_fin.empty:
                    st.markdown("### 💳 Financiamientos")
                    df_fin["Saldo Pendiente"]=pd.to_numeric(df_fin["Saldo Pendiente"],errors="coerce").fillna(0)
                    df_fin["Monto Financiado"]=pd.to_numeric(df_fin["Monto Financiado"],errors="coerce").fillna(0)
                    act=df_fin[df_fin["Estado Fin"].astype(str).str.lower()=="activo"]
                    rf=st.columns(3)
                    with rf[0]: st.metric("Activos",len(act))
                    with rf[1]: st.metric("Total financiado",fmt(df_fin["Monto Financiado"].sum()))
                    with rf[2]: st.metric("Saldo pendiente",fmt(act["Saldo Pendiente"].sum()))
                    st.dataframe(df_fin,use_container_width=True)

                st.divider()
                st.markdown("### 🚗 Costo y ganancia por vehículo")
                cols_inf=[c for c in ["ID","Vehículo","Año","Estado","Días Inventario","Costo Total","Precio Objetivo (CRC)","Precio Venta","Ganancia Real","Ganancia Edgardo","Ganancia Socio Real"] if c in df_calc.columns]
                st.dataframe(df_calc[cols_inf],use_container_width=True)

                st.divider()
                if not df_ven.empty:
                    st.markdown("### 📅 Ventas por mes")
                    dv=df_ven.copy(); dv["Fecha Venta"]=pd.to_datetime(dv["Fecha Venta"],errors="coerce")
                    dv["Mes"]=dv["Fecha Venta"].dt.strftime("%Y-%m"); dv["Precio CRC"]=pd.to_numeric(dv["Precio CRC"],errors="coerce").fillna(0)
                    vm=dv.groupby("Mes").agg(Total=("Precio CRC","sum"),Cantidad=("Precio CRC","count")).reset_index()
                    st.dataframe(vm,use_container_width=True)

                st.divider()
                if not df_op.empty:
                    st.markdown("### 🏢 Gastos operativos por tipo")
                    df_op["Monto CRC"]=pd.to_numeric(df_op["Monto CRC"],errors="coerce").fillna(0)
                    gt=df_op.groupby("Tipo")["Monto CRC"].sum().reset_index().sort_values("Monto CRC",ascending=False)
                    fig_op=px.bar(gt,x="Tipo",y="Monto CRC",title="Gastos por tipo",color_discrete_sequence=["#e63946"])
                    fig_op.update_layout(paper_bgcolor="#1a1a1a",plot_bgcolor="#1a1a1a",font_color="#f5f0e8")
                    st.plotly_chart(fig_op,use_container_width=True)

                if not df_ven.empty and "Vendedor" in df_ven.columns:
                    st.divider(); st.markdown("### 👥 Comisiones por vendedor")
                    df_ven["Comisión CRC"]=pd.to_numeric(df_ven["Comisión CRC"],errors="coerce").fillna(0)
                    df_ven["Precio CRC"]=pd.to_numeric(df_ven["Precio CRC"],errors="coerce").fillna(0)
                    com=df_ven.groupby("Vendedor").agg(Comisiones=("Comisión CRC","sum"),Ventas=("Precio CRC","sum"),Cantidad=("Precio CRC","count")).reset_index()
                    st.dataframe(com,use_container_width=True)

                if "Días Inventario" in df_calc.columns:
                    st.divider(); st.markdown("### ⏱️ Carros +60 días en inventario")
                    lentos=df_calc[(df_calc["Días Inventario"]>60)&(df_calc["Estado"].astype(str).str.lower()=="disponible")][["ID","Vehículo","Año","Días Inventario","Costo Total","Precio Objetivo (CRC)"]]
                    if not lentos.empty: st.dataframe(lentos,use_container_width=True)
                    else: st.success("✅ Todos los carros tienen menos de 60 días.")

        # ══ NUEVO VEHÍCULO ══
        with tabs[11]:
            es_canje=st.session_state.get("ir_a_canje",False)
            if es_canje: st.info("🔄 Registrando vehículo recibido como CANJE")
            st.markdown("### ➕ Agregar nuevo vehículo")
            with st.form("form_nuevo"):
                c1,c2,c3=st.columns(3)
                id_n=c1.number_input("ID",min_value=1,step=1,value=len(df_inv)+1 if not df_inv.empty else 1)
                vehiculo_n=c2.text_input("Vehículo (marca y modelo)")
                año_n=c3.number_input("Año",min_value=1950,max_value=2030,step=1,value=2020)
                c4,c5,c6=st.columns(3)
                color_n=c4.text_input("Color"); placa_n=c5.text_input("Placa"); vin_n=c6.text_input("VIN")
                c7,c8,c9=st.columns(3)
                socios_n=c7.text_input("Socios")
                tipo_prop_n=c8.selectbox("Tipo Propiedad",TIPOS_PROPIEDAD)
                origen_n=c9.selectbox("Origen",["Compra","Canje","Recibo"],index=1 if es_canje else 0)
                c10,c11,c12=st.columns(3)
                pct_g=c10.number_input("% Grupo (0-1)",min_value=0.0,max_value=1.0,step=0.05,value=0.5 if tipo_prop_n=="Con Socio" else 1.0)
                pct_s=c11.number_input("% Socio (0-1)",min_value=0.0,max_value=1.0,step=0.05,value=0.5 if tipo_prop_n=="Con Socio" else 0.0)
                precio_obj_n=c12.number_input("Precio Objetivo (CRC)",min_value=0.0,step=100000.0)
                c13,c14=st.columns(2)
                precio_min_n=c13.number_input("Precio Mínimo (CRC)",min_value=0.0,step=100000.0)
                url_foto_n=c14.text_input("URL de foto (opcional)")
                fecha_entrada_n=st.date_input("Fecha entrada al inventario",value=date.today())
                notas_n=st.text_area("Notas")
                submit_n=st.form_submit_button("✅ Agregar vehículo")
            if submit_n:
                try:
                    get_ws("INVENTARIO").append_row([int(id_n),vehiculo_n,int(año_n),color_n,placa_n,vin_n,socios_n,tipo_prop_n,pct_g,pct_s,"Disponible",precio_obj_n,precio_min_n,str(fecha_entrada_n),notas_n,origen_n,url_foto_n])
                    if es_canje: st.session_state["ir_a_canje"]=False
                    invalidar(); st.success(f"✅ {vehiculo_n} agregado."); st.rerun()
                except Exception as e: st.error(f"Error: {e}")

        # ══ CONFIG ══
        with tabs[12]:
            st.markdown("### ⚙️ Configuración")
            st.markdown("#### 👥 Vendedores")
            with st.form("form_vend"):
                id_v_n=(len(df_vend)+1) if not df_vend.empty else 1
                c1,c2=st.columns(2)
                nv=c1.text_input("Nombre"); tel_v=c2.text_input("Teléfono")
                sv=st.form_submit_button("Agregar vendedor")
            if sv and nv:
                try:
                    get_ws("VENDEDORES").append_row([id_v_n,nv,tel_v,"Sí"])
                    invalidar(); st.success(f"✅ Vendedor {nv} agregado."); st.rerun()
                except Exception as e: st.error(f"Error: {e}")
            if not df_vend.empty: st.dataframe(df_vend,use_container_width=True)

        # ══ IMPORTAR ══
        with tabs[13]:
            st.markdown("### 📥 Importar desde Excel")
            st.info("Sube la plantilla completa. Se importarán todas las hojas automáticamente.")
            archivo=st.file_uploader("Subir Excel",type=["xlsx"])
            if archivo:
                try:
                    xls=pd.ExcelFile(archivo)
                    for hoja in xls.sheet_names:
                        if hoja in HOJAS_HEADERS:
                            df_prev=pd.read_excel(xls,sheet_name=hoja).dropna(how='all')
                            if not df_prev.empty:
                                st.markdown(f"**{hoja}:** {len(df_prev)} filas")
                                st.dataframe(df_prev.head(3),use_container_width=True)
                    if st.button("✅ Importar TODAS las hojas"):
                        total=0; errores=[]
                        for hoja in xls.sheet_names:
                            if hoja in HOJAS_HEADERS:
                                try:
                                    df_xl=pd.read_excel(xls,sheet_name=hoja).dropna(how='all')
                                    primera_col=HOJAS_HEADERS[hoja][0]
                                    if primera_col in df_xl.columns:
                                        df_xl=df_xl[df_xl[primera_col].notna()]
                                        df_xl=df_xl[df_xl[primera_col].astype(str).str.lower()!='none']
                                    if df_xl.empty: continue
                                    cols_imp=HOJAS_HEADERS[hoja]
                                    ws_imp=get_ws(hoja)
                                    datos_act=ws_imp.get_all_values()
                                    if not datos_act or datos_act[0]!=cols_imp:
                                        ws_imp.clear(); ws_imp.append_row(cols_imp)
                                    imp=0
                                    for _,row in df_xl.iterrows():
                                        fila=[]
                                        for c in cols_imp:
                                            val=row.get(c,"")
                                            if pd.isna(val) or str(val).lower()=='none': fila.append("")
                                            else:
                                                vs=str(val).strip()
                                                if vs=='COL': vs='CRC'
                                                if vs.endswith('.0') and vs[:-2].isdigit(): vs=vs[:-2]
                                                fila.append(vs)
                                        ws_imp.append_row(fila); imp+=1
                                    total+=imp; st.success(f"✅ {hoja}: {imp} filas")
                                except Exception as e: errores.append(f"{hoja}: {e}")
                        for err in errores: st.error(f"❌ {err}")
                        invalidar(); st.success(f"🎉 Total: {total} filas importadas."); st.rerun()
                except Exception as e: st.error(f"Error: {e}")

        # ══ EXPORTAR ══
        with tabs[14]:
            st.markdown("### 📤 Exportar todos los datos")
            if st.button("📥 Generar Excel completo"):
                try:
                    wb_exp=Workbook(); primera=True
                    hojas_exp={"INVENTARIO":df_calc,"GASTOS_VEHICULO":df_gv,"VENTAS":df_ven,"CANJES":df_canjes,"FINANCIAMIENTOS":df_fin,"PAGOS":df_pagos,"GASTOS_OPERATIVOS":df_op,"CLIENTES":df_cli,"CONSIGNACIONES":df_cons}
                    for nombre,df_e in hojas_exp.items():
                        if primera: ws_e=wb_exp.active; ws_e.title=nombre; primera=False
                        else: ws_e=wb_exp.create_sheet(nombre)
                        if not df_e.empty:
                            for j,col in enumerate(df_e.columns,1):
                                c=ws_e.cell(row=1,column=j,value=col)
                                c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",start_color="C0392B")
                                c.alignment=Alignment(horizontal="center")
                            for i,(_,row) in enumerate(df_e.iterrows(),2):
                                for j,val in enumerate(row,1):
                                    ws_e.cell(row=i,column=j,value=str(val) if val else "")
                            for col in ws_e.columns:
                                ws_e.column_dimensions[get_column_letter(col[0].column)].width=18
                    buf=io.BytesIO(); wb_exp.save(buf)
                    st.download_button("⬇️ Descargar Excel",data=buf.getvalue(),file_name=f"GalleryMotors_{date.today()}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                except Exception as e: st.error(f"Error: {e}")

        # ══ EMAIL ══
        with tabs[15]:
            st.markdown("### 📧 Enviar informe por correo")
            import smtplib
            from email.mime.multipart import MIMEMultipart
            from email.mime.text import MIMEText
            st.info("Se enviará a: edgardo@royzacr.com y gerlad@royzacr.com")
            periodo_e=st.selectbox("Período",["Semanal","Mensual","General"])
            col_e1,col_e2=st.columns(2)
            with col_e1:
                if st.button("📧 Enviar informe"):
                    try:
                        gan_e=df_calc["Ganancia Edgardo"].sum() if not df_calc.empty and "Ganancia Edgardo" in df_calc.columns else 0
                        total_i=df_calc["Costo Total"].sum() if not df_calc.empty else 0
                        gastos_o=pd.to_numeric(df_op["Monto CRC"],errors="coerce").sum() if not df_op.empty and "Monto CRC" in df_op.columns else 0
                        html_e=f"""<html><body style="font-family:Arial;max-width:700px;margin:auto;">
                        <div style="background:#0a0a0a;padding:20px;text-align:center;border-bottom:3px solid #e63946;">
                        <h1 style="color:#e63946;margin:0;">Gallery Motors by UR</h1>
                        <p style="color:#888;">Informe {periodo_e} — {datetime.now().strftime('%d/%m/%Y')}</p></div>
                        <div style="padding:20px;">
                        <table style="width:100%;border-collapse:collapse;">
                        <tr style="background:#f5f5f5;"><td style="padding:10px;">Total Invertido</td><td style="padding:10px;font-weight:bold;">{fmt(total_i)}</td></tr>
                        <tr><td style="padding:10px;">Mi Ganancia</td><td style="padding:10px;font-weight:bold;color:green;">{fmt(gan_e)}</td></tr>
                        <tr style="background:#f5f5f5;"><td style="padding:10px;">Gastos Operativos</td><td style="padding:10px;font-weight:bold;color:red;">{fmt(gastos_o)}</td></tr>
                        <tr><td style="padding:10px;">Disponibles</td><td style="padding:10px;font-weight:bold;">{len(df_calc[df_calc["Estado"].astype(str).str.lower()=="disponible"]) if not df_calc.empty else 0}</td></tr>
                        </table></div></body></html>"""
                        msg=MIMEMultipart('alternative')
                        msg['Subject']=f"Gallery Motors — Informe {periodo_e} {datetime.now().strftime('%d/%m/%Y')}"
                        msg['From']="edgardo@royzacr.com"; msg['To']="edgardo@royzacr.com, gerlad@royzacr.com"
                        msg.attach(MIMEText(html_e,'html'))
                        with smtplib.SMTP_SSL('smtp.gmail.com',465) as server:
                            server.login("edgardo@royzacr.com","rcwd pvcy kuxe qewc")
                            server.sendmail("edgardo@royzacr.com",["edgardo@royzacr.com","gerlad@royzacr.com"],msg.as_string())
                        st.success("✅ Informe enviado!")
                    except Exception as e: st.error(f"Error: {e}")
            with col_e2:
                gan_e2=df_calc["Ganancia Edgardo"].sum() if not df_calc.empty and "Ganancia Edgardo" in df_calc.columns else 0
                total_i2=df_calc["Costo Total"].sum() if not df_calc.empty else 0
                gastos_o2=pd.to_numeric(df_op["Monto CRC"],errors="coerce").sum() if not df_op.empty and "Monto CRC" in df_op.columns else 0
                html_dl=f"<html><body><h1>Gallery Motors</h1><p>Total: {fmt(total_i2)}</p><p>Ganancia: {fmt(gan_e2)}</p><p>Gastos: {fmt(gastos_o2)}</p></body></html>"
                st.download_button("⬇️ Descargar informe",data=html_dl,file_name=f"Informe_{date.today()}.html",mime="text/html")

        # ══ CITAS ══
        with tabs[16]:
            st.markdown("### 📅 Registrar cita")
            with st.form("form_cita"):
                c1,c2,c3=st.columns(3)
                id_cita_v=c1.selectbox("Vehículo de interés",options=ids_v if ids_v else ["Sin vehículos"],format_func=vlabel if ids_v else lambda x:x)
                fecha_cita=c2.date_input("Fecha",value=date.today())
                hora_cita=c3.time_input("Hora",value=datetime.now().replace(minute=0,second=0).time())
                c4,c5=st.columns(2)
                id_cita_cli=c4.selectbox("Cliente",options=["Sin cliente"]+ids_cli,format_func=lambda x: clabel(x) if x!="Sin cliente" else "Sin cliente")
                vendedor_cita=c5.selectbox("Vendedor asignado",options=["Sin asignar"]+vend_lista)
                notas_cita=st.text_area("Notas")
                id_cita_n=(len(df_citas)+1) if not df_citas.empty else 1
                submit_cita=st.form_submit_button("💾 Registrar cita")
            if submit_cita:
                try:
                    get_ws("CITAS").append_row([id_cita_n,str(id_cita_v),str(id_cita_cli),str(fecha_cita),str(hora_cita),vendedor_cita,"Pendiente",notas_cita,""])
                    # WhatsApp link para vendedor
                    if vendedor_cita!="Sin asignar" and not df_vend.empty:
                        vr=df_vend[df_vend["Nombre"]==vendedor_cita]
                        if not vr.empty:
                            tel=str(vr.iloc[0].get("Teléfono",""))
                            if tel:
                                tel=tel.replace("-","").replace(" ","")
                                if not tel.startswith("506"): tel="506"+tel
                                msg_wa=f"🚗 Nueva cita Gallery Motors\nVehículo: {vlabel(id_cita_v)}\nFecha: {fecha_cita}\nHora: {hora_cita}"
                                link_wa=f"https://wa.me/{tel}?text={urllib.parse.quote(msg_wa)}"
                                invalidar()
                                st.success("✅ Cita registrada.")
                                st.markdown(f'<a href="{link_wa}" target="_blank" style="display:inline-block;background:#25D366;color:white;padding:0.5rem 1rem;border-radius:4px;text-decoration:none;font-weight:600;margin-top:0.5rem;">📲 Notificar a {vendedor_cita} por WhatsApp</a>',unsafe_allow_html=True)
                                st.rerun()
                    invalidar(); st.success("✅ Cita registrada."); st.rerun()
                except Exception as e: st.error(f"Error: {e}")

            if not df_citas.empty:
                st.divider()
                st.markdown("### Citas registradas")
                filtro_c=st.selectbox("Filtrar",["Todas","Pendiente","Completada","Cancelada"])
                df_c_s=df_citas.copy()
                if filtro_c!="Todas": df_c_s=df_c_s[df_c_s["Estado"].astype(str)==filtro_c]
                st.dataframe(df_c_s,use_container_width=True)
                st.divider()
                st.markdown("### ✏️ Actualizar cita")
                if not df_citas.empty and "ID" in df_citas.columns:
                    with st.form("form_upd_cita"):
                        id_c_upd=st.selectbox("Cita",options=df_citas["ID"].astype(str).tolist())
                        c1,c2=st.columns(2)
                        nuevo_est_c=c1.selectbox("Estado",["Pendiente","Completada","Cancelada"])
                        seguimiento=c2.text_input("Notas de seguimiento")
                        submit_upd_c=st.form_submit_button("💾 Actualizar")
                    if submit_upd_c:
                        try:
                            ws_c=get_ws("CITAS"); data_c=ws_c.get_all_values()
                            for idx,fila in enumerate(data_c):
                                if idx==0: continue
                                if str(fila[0])==str(id_c_upd):
                                    ws_c.update_cell(idx+1,data_c[0].index("Estado")+1,nuevo_est_c)
                                    ws_c.update_cell(idx+1,data_c[0].index("Seguimiento")+1,seguimiento); break
                            invalidar(); st.success(f"✅ Cita actualizada."); st.rerun()
                        except Exception as e: st.error(f"Error: {e}")

        # ══ LOG ══
        with tabs[17]:
            st.markdown("### 📋 Log de actividad")
            with st.form("form_log"):
                c1,c2,c3=st.columns(3)
                id_log_v=c1.selectbox("Vehículo",options=ids_v if ids_v else ["Sin vehículos"],format_func=vlabel if ids_v else lambda x:x)
                tipo_log=c2.selectbox("Tipo",["Nota","Visita","Llamada","Precio cambiado","Reparación","Otro"])
                fecha_log=c3.date_input("Fecha",value=date.today())
                desc_log=st.text_area("Descripción")
                usuario_log=st.text_input("Registrado por",value="Edgardo")
                submit_log=st.form_submit_button("💾 Guardar")
            if submit_log:
                try:
                    get_ws("LOG_ACTIVIDAD").append_row([str(id_log_v),str(fecha_log),tipo_log,desc_log,usuario_log])
                    invalidar(); st.success("✅ Evento registrado."); st.rerun()
                except Exception as e: st.error(f"Error: {e}")
            if not df_log.empty and ids_v:
                st.divider()
                id_log_sel=st.selectbox("Ver log de",options=ids_v,format_func=vlabel)
                lf=df_log[df_log["ID Vehículo"].astype(str)==str(id_log_sel)]
                if not lf.empty: st.dataframe(lf,use_container_width=True)
                else: st.info("Sin actividad.")

